"""Generate submission/NotesNPU_Description.docx (and .pdf via Word when --pdf is passed).

python submission/build_description.py --pdf
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402

RED = RGBColor(0xE4, 0x1E, 0x2B)
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def h(text, level=1):
    p = doc.add_heading(text, level)
    for r in p.runs:
        r.font.color.rgb = RED if level == 1 else RGBColor(0x14, 0x17, 0x1F)
    return p


def bl(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def tbl(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = widths[j]
            c.text = str(val)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.bold = i == 0
                    if i == 0:
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                shade(c, "E41E2B")
    doc.add_paragraph()


title = doc.add_paragraph()
r = title.add_run(C.TITLE)
r.font.size, r.font.bold, r.font.color.rgb = Pt(28), True, RED
sub = doc.add_paragraph()
r = sub.add_run(C.TAGLINE)
r.font.size, r.font.italic = Pt(13), True
meta = doc.add_paragraph()
meta.add_run(f"Snapdragon AI Lab Build & Present Challenge 2026  |  {C.AUTHOR}  |  {C.REPO}").font.size = Pt(9.5)

h("1. Overview")
doc.add_paragraph(
    "NotesNPU is a Windows desktop application for Snapdragon-powered HP PCs. It turns lectures, classes and "
    "meetings into timestamped transcripts, summaries, structured study notes, flashcards, quizzes and a "
    "grounded Q&A chatbot, completely on-device. Speech recognition runs Qualcomm AI Hub's Whisper model "
    "on the Snapdragon Hexagon NPU through the ONNX Runtime QNN Execution Provider. A local small language "
    "model (Phi-3.5-mini, INT4) produces the study material. No audio or text ever leaves the laptop, and "
    "it works without internet.")

h("2. Problem")
bl(C.PROBLEM)

h("3. Solution & key features")
bl(C.SOLUTION)
tbl([("Feature", "What it does"), *C.FEATURES], [Inches(1.6), Inches(5.3)])

h("4. AI models used")
tbl([("Task", "Model", "Source", "Runtime / hardware"), *C.MODELS],
    [Inches(1.5), Inches(2.0), Inches(1.5), Inches(2.0)])

h("5. Technical implementation")
bl([
    "Audio pipeline: 16 kHz mono resampling, pause-aware chunking (<=30 s windows cut at the quietest 100 ms), RMS silence gate to prevent Whisper hallucinations.",
    "NPU inference: AI Hub 'precompiled QNN ONNX' Whisper-Base for Snapdragon X Elite / X2 Elite (FP16 EPContext graphs, QAIRT 2.50). "
    "The encoder outputs cross-attention K/V caches directly. A fixed-shape single-step decoder (200-slot self-attention KV cache, "
    "additive attention mask, position ids) keeps every step on the HTP. The QNN plugin EP is registered at runtime "
    "(onnxruntime-qnn 2.x) with the HTP backend and burst performance mode.",
    "Decoding: forced prompt <|sot|><|lang|><|task|><|notimestamps|>, greedy decoding with timestamp suppression, n-gram repetition guard, "
    "automatic language detection, and translate-to-English mode for Indian languages.",
    "Study AI: onnxruntime-genai streaming generation with task-specific prompts, head+tail context fitting for long lectures, "
    "Markdown clean-up, and an instant extractive fallback (TF-IDF ranking, action-item detection, cloze MCQs).",
    "Robustness: automatic NPU to CPU to extractive fallback, so the same build runs on any Windows PC; built-in NPU vs CPU benchmark.",
    "UI: local Gradio web app (127.0.0.1 only) with Transcribe, Live mode, Study notes, Q&A and Benchmark tabs, plus Markdown export.",
])

h("6. Why Snapdragon / NPU optimisation")
bl(C.WHY_SNAPDRAGON)
tbl([("Device / backend", "Speed", "Encoder / 30 s", "Decoder / token"), *C.BENCH_ROWS],
    [Inches(2.9), Inches(1.4), Inches(1.3), Inches(1.3)])

h("7. Use cases & impact")
tbl([("Who", "How NotesNPU helps"), *C.IMPACT], [Inches(1.8), Inches(5.1)])

h("8. Deployment & accessibility")
bl(C.DEPLOYMENT)

h("9. Roadmap")
bl(C.ROADMAP)

h("10. Originality & licences")
doc.add_paragraph(C.ORIGINALITY)

shot = ROOT / "docs" / "screenshots" / "notes.png"
if not shot.exists():
    shot = ROOT / "docs" / "screenshots" / "transcribe.png"
if shot.exists():
    h("Appendix: screenshot")
    doc.add_picture(str(shot), width=Inches(4.2))

out = HERE / "NotesNPU_Description.docx"
doc.save(out)
print("wrote", out)

if "--pdf" in sys.argv:
    try:
        import comtypes.client  # type: ignore

        word = comtypes.client.CreateObject("Word.Application")
        d = word.Documents.Open(str(out))
        d.SaveAs(str(out.with_suffix(".pdf")), FileFormat=17)
        d.Close()
        word.Quit()
        print("wrote", out.with_suffix(".pdf"))
    except Exception as e:
        print(f"(PDF export skipped: {e}) - open the DOCX in Word > Save As > PDF")
